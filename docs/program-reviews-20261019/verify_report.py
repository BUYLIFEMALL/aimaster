"""Validate the report's coverage, evidence, classification, links and Excel independently."""
import collections,json,pathlib,re,subprocess,urllib.parse
from openpyxl import load_workbook
OUT=pathlib.Path(__file__).resolve().parent;ROOT=OUT.parents[1]
d=json.loads((OUT/'review-data.json').read_text(encoding='utf-8'));checks=[]
def check(name,ok):
 if not ok:raise AssertionError(name)
 checks.append({'name':name,'passed':True})
p=d['programs'];items=[i for a in p for i in a['items']]
check('83 unique programs and complete catalog coverage',len(p)==83 and len({x['slug']for x in p})==83 and {x['slug']for x in p}=={x['slug']for x in json.loads((OUT/'review-basis.json').read_text(encoding='utf-8'))['programs']})
check('265 unique detailed items',len(items)==265 and len({i['id']for i in items})==265)
check('14 categories',len({a['category']for a in p})==14)
check('Every program has all three perspectives',all({i['perspective']for i in a['items']}=={'개발·유지보수','디자인·사용성','신규 자동화'}for a in p))
fields=['current','why','proposal','acceptance','confirmation','stability','applicability','convenience','status']
check('Every item has a reason, evidence limits, scope, acceptance and adoption criteria',all(all(isinstance(i[k],str)and len(i[k])>=10 for k in fields)for i in items))
check('Items are program-specific, no repeated reason/template rows',len({i['why']for i in items})==265 and len({i['acceptance']for i in items})==265)
ev=json.loads((OUT/'source-evidence.json').read_text(encoding='utf-8'));by={x['path']:x for x in ev}
check('All evidence locations exist at the stated revision',all(i['evidence_file']in by and 0<i['evidence_line']<=by[i['evidence_file']]['lines']and by[i['evidence_file']]['revision']==d['source_revision']for i in items))
check('All evidence belongs to the reviewed program',all(i['evidence_file'].startswith(a['folder']+'/')for a in p for i in a['items']))
check('Frozen source links preserve paths including bracket routes',all(urllib.parse.unquote(i['evidence_url'].split('/blob/'+d['source_revision']+'/')[1].split('#')[0])==i['evidence_file']for i in items))
for a in p:
 md=(OUT/'programs'/f"{a['slug']}.md").read_text(encoding='utf-8')
 check('Dossier complete: '+a['slug'],all('## '+i['id']+' ·'in md and i['why']in md and i['acceptance']in md for i in a['items']) and a['app_url']in md)
logs=[json.loads(l)for l in (OUT/'read-log.jsonl').read_text(encoding='utf-8').splitlines()if l.strip()]
check('Each program has selected core-source reading evidence',all(any(l['file'].startswith(a['folder']+'/')and re.search(r'\.(ts|tsx|js|jsx|mjs|py)$',l['file'])for l in logs)for a in p))
wb=load_workbook(OUT/'AIMaster_프로그램별_상세개선_20261019.xlsx',data_only=False)
check('Excel 7 sheets and exact program/item counts',len(wb.sheetnames)==7 and wb['프로그램83'].max_row==84 and wb['상세개선265'].max_row==266 and wb['채택기준'].max_row==266)
check('Excel rows preserve detailed reasons and acceptance',all(wb['상세개선265'].cell(n,7).value==i['why']and wb['상세개선265'].cell(n,9).value==i['acceptance']for n,i in enumerate(items,2)))
check('Excel has filters and frozen headings',all(ws.auto_filter.ref and ws.freeze_panes=='A2'for ws in wb))
check('Excel no executable formulas',not any(c.data_type=='f'for ws in wb for row in ws for c in row))
page=(OUT/'REPORT.html').read_text(encoding='utf-8')
check('HTML exact 83 program cards and 265 item sections',page.count('<article class="program"')==83 and page.count('<section class="item"')==265)
check('HTML no external scripts/fonts or implicit product requests',not re.search(r'<(?:script|link|img)[^>]+(?:src|href)=[\"\']https?://',page)and not re.search(r'\bfetch\(',page))
check('Readability: no long unspaced Korean fragments',not any(re.search(r'[가-힣]{14,}',i[k])for i in items for k in ['current','why','proposal','acceptance']))
pub=OUT/'publication-manifest.json'
if pub.exists():
 manifest=json.loads(pub.read_text(encoding='utf-8'));rows=manifest['programs']
 check('All 83 folder publication decisions recorded',len(rows)==83 and {r['slug']for r in rows}=={a['slug']for a in p})
 check('Published folder documents exist and link to central snapshot',all((ROOT/r['document']).exists()and 'program-reviews-20261019'in (ROOT/r['document']).read_text(encoding='utf-8')for r in rows if r['state']=='installed'))
 check('Deferred folders have explicit reasons and full central dossiers',all(r.get('reason')and(OUT/'programs'/f"{r['slug']}.md").exists()for r in rows if r['state']=='deferred'))
browser=OUT/'report-browser-verification.json'
if browser.exists():
 b=json.loads(browser.read_text(encoding='utf-8'));check('Headed Microsoft Edge report checks passed, no external requests',b['browser'].startswith('Microsoft Edge')and len(b['checks'])>=12 and all(c['passed']for c in b['checks'])and not b['errors']and not b['external_network'])
(OUT/'artifact-verification.json').write_text(json.dumps({'scope':'Report artifacts only; not product feature tests','checks':checks,'passed':len(checks),'programs':len(p),'items':len(items)},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'passed':len(checks),'programs':len(p),'items':len(items)},ensure_ascii=False))
