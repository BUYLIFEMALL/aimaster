"""Render the analyst-authored review snapshot. Does not call product APIs or write app folders."""
from __future__ import annotations
import collections,hashlib,html,json,pathlib,re,runpy,subprocess,urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter

OUT=pathlib.Path(__file__).resolve().parent
ROOT=OUT.parents[1]
def git(*args):return subprocess.check_output(['git',*args],cwd=ROOT).decode('utf-8-sig',errors='replace')
def write(name,value):
 p=OUT/name;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(value,encoding='utf-8',newline='\n')
def dump(name,value):write(name,json.dumps(value,ensure_ascii=False,indent=2)+'\n')
def esc(s):return html.escape(str(s),quote=True)
def cell(s):return str(s).replace('|','&#124;').replace('\n','<br>')
def table(headers,rows):return '| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(cell(v) for v in r)+' |\n' for r in rows)

basis=json.loads((OUT/'review-basis.json').read_text(encoding='utf-8'))
BASE=basis['source_revision'];HEAD=git('rev-parse','HEAD').strip()
prior=json.loads((ROOT/'docs/analysis-20261019/proposal-data.json').read_text(encoding='utf-8'))
prior_apps={x['slug']:x for x in prior['programs']}
raw={}
for file in [OUT/'reviews.py',*sorted(OUT.glob('batch_*.py'))]:
 d=runpy.run_path(str(file))['REVIEWS'];assert not raw.keys()&d.keys();raw.update(d)
assert len(raw)==83 and len(basis['programs'])==83
prose_fixes={'크 롤 링':'크롤링','내 보내기':'내보내기','내 보내는':'내보내는','로그 아웃':'로그아웃','실행 기는':'실행기는','랜드 마크':'랜드마크','원 생성 일':'원래 생성일','원생성일':'원래 생성일','캐러셀게시':'캐러셀 게시','캐러셀완료개수':'캐러셀 완료 개수','중복태그를로컬':'중복 태그를 로컬','Config로형 변환':'Config로 형 변환','한다.300':'한다. 300','뼈대복사':'뼈대 복사','동시수':'동시 수','안전 내 보내기':'안전 내보내기','최저가대표값':'최저가 대표값','query 해시로공용조회':'query 해시로 공용 조회'}
def polish(s):
 for a,b in prose_fixes.items():s=s.replace(a,b)
 for a,b in {
 '지연예약을가짜제공사로검사하며수집':'지연 예약을 가짜 제공자로 검사하며 수집',
 '부분캐러셀실패에서기존게시':'캐러셀 일부 실패에서 기존 게시',
 '메시지별답장을별도멱등성단계로두고외부':'메시지별 답장을 별도 중복 방지 단계로 두고 외부',
 '모바일스와이프가기존후보며카드수':'모바일 스와이프가 기존 후보이며 카드 수',
 '타래초안비교와수동내보내기를':'타래 초안 비교와 수동 내보내기를',
 '대기로보고다시결제성렌더를':'대기로 보고 다시 유료 렌더링을',
 '세로가산휴리스틱이고목표길이보정':'세로 가산 휴리스틱이고 목표 길이 보정',
 '소유자채널검증을두되웹훅에일반로그인':'소유자·채널 검증을 두되 웹훅에 일반 로그인',
 '상품변수를프롬프트템플릿':'상품 변수를 프롬프트 템플릿',
 '게시성공뒤기록실패를':'게시 성공 뒤 기록 실패를','관리자웹훅이전은':'관리자 웹훅 이전은',
 '독립운영키트빌드는':'독립 운영 키트 빌드는','빈자산만채우기는':'빈 자산만 채우기는',
 '순위점수금지프롬프트는':'순위 점수 금지 프롬프트는','실패건다시보내기는':'실패 건 다시 보내기는',
 '알림톡비친구도달은':'알림톡의 비친구 도달은','앱미등록오류배너누락은':'앱 미등록 오류 배너 누락은',
 '에긴개인프롬프트를':'에 긴 개인 프롬프트를','여부조회후외부게시하고':'여부 조회 후 외부 게시하고',
 '와원자적이력추가':'와 원자적 이력 추가','완료항목재게시가':'완료 항목 재게시가',
 '응답유실에서렌더를':'응답 유실에서 렌더링을','있음두표본에서허위구매':'있음 두 표본에서 허위 구매',
 '조회후외부답글을':'조회 후 외부 답글을','종료된배너오류를':'종료된 배너 오류를',
 '컷설명이프롬프트에':'컷 설명이 프롬프트에','키없음뒤프롬프트를':'키 없음 뒤 프롬프트를',
 '회원지시와대화보관이':'회원 지시와 대화 보관이','구분해야한다':'구분해야 한다',
 '입력해 시':'입력 해시','형 변환':'형식 단언','실행 ID로':'실행 ID로',
 '기존 제안 상세화':'기존 제안을 구체화한 검토 후보','F-012 재확인':'F-012의 자동 제공자 전환 코드 재확인',
 'I-14 상세화':'I-14 비교 작업실 제안의 프로그램별 구체화','반환 분기 확인':'실행 결과의 반환 분기 코드 확인',
 '폴백 분기 확인':'자동 유형 판별 실패 후 대체 분기 확인'
 }.items():s=s.replace(a,b)
 return re.sub(r'([.!?])(?=[가-힣])',r'\1 ',s)

programs=[];evidence={}
for n,p in enumerate(basis['programs'],1):
 slug=p['slug'];folder=p['folder'];items=[]
 for j,r in enumerate(raw[slug],1):
  assert len(r)==9
  priority,view,title,ref,current,why,proposal,acceptance,confirmation=r
  if slug=='ebook-studio' and j==1:priority='P0'
  rel,line=ref.rsplit(':',1);line=int(line);file=folder+'/'+rel
  if file not in evidence:
   src=git('show',BASE+':'+file)
   evidence[file]={'path':file,'revision':BASE,'sha256_utf8':hashlib.sha256(src.encode()).hexdigest(),'lines':len(src.splitlines())}
  assert 1<=line<=evidence[file]['lines'],(slug,ref)
  current,why,proposal,acceptance,confirmation=map(polish,[current,why,proposal,acceptance,confirmation])
  # Existing strengths and regression tasks are not certified production bugs.
  kind='기능 확장 후보' if view=='신규 자동화' else '화면·설명 개선 후보' if view=='디자인·사용성' else '코드 경계·회귀 검증 과제'
  items.append({'id':f'D-{n:02d}-{j:02d}','priority':priority,'perspective':view,'title':polish(title),'kind':kind,'evidence_file':file,'evidence_line':line,'evidence_url':'https://github.com/BUYLIFEMALL/aimaster/blob/'+BASE+'/'+urllib.parse.quote(file,safe='/')+'#L'+str(line),'current':current,'why':why,'proposal':proposal,'acceptance':acceptance,'confirmation':confirmation,'status':'검토 후보 — 구현 상태는 최신 프로그램 원장과 대조','stability':'검증 대기: '+acceptance,'applicability':'대상 근거: '+ref+'; 최신 코드·회원 권한을 재대조한 뒤 한 프로그램씩 적용','convenience':why,'implementation_boundary':'개인 계정·본인 API·소유자별 자료; 생성/조사 자료는 최초 생성 +30일. 분석은 구현 승인이 아님.','approval_note':'DB/환경변수/키/파괴적 삭제/대량 유료 호출이 포함되면 기존 승인 절차 적용. 그 외도 작업 전 개발 플랜 제시.'})
 version=[]
 for v in p.get('code_versions',[]):
  try:
   text=git('show',BASE+':'+v['file']);m=re.search(r'APP_VERSION\s*=\s*[\'\"]([^\'\"]+)',text)
   if m:version.append(m[1])
  except subprocess.CalledProcessError:pass
 previous=prior_apps.get(slug,{})
 programs.append({'sequence':n,'slug':slug,'name':p['name'],'folder':folder,'category':p['category'],'app_url':p['app_url'],'catalog_version_previous_snapshot':p['version'],'code_version_at_basis':', '.join(version) or '개별 버전 파일 재확인','short_desc':p.get('short_desc',''),'workflow':p.get('baseline_capabilities',{}).get('기능') or p.get('short_desc',''),'strength':previous.get('strength') or p.get('baseline_capabilities',{}).get('장점') or '기존 동작과 회원별 계정 경계를 보존하며 개선한다.','items':items})

drift=git('diff','--name-only',BASE,HEAD,'--',*evidence).splitlines()
dump('source-drift.json',{'basis':BASE,'compared_head':HEAD,'changed_evidence_files':drift,'changed_core_files':[f for f in drift if re.search(r'\.(?:ts|tsx|js|jsx|py)$',f)],'note':'기준 이후 변경 파일 목록. 카드뉴스 문구 리메이크 분기와 GIF 화면 변경은 LATEST_RECHECK.md에서 재대조 결과 구분. 구현 전에 최신 원장·ACTIVE_WORK 재확인.'})
dump('source-evidence.json',list(evidence.values()))
data={'date':'2026-10-19','source_revision':BASE,'comparison_revision':HEAD,'scope':'83개 카탈로그 프로그램의 문서와 선택 핵심 소스 경로에 대한 정적 분석. 전체 소스 전수 정독·83개 운영 기능 검증·보안 인증이 아님.','counts':{'programs':len(programs),'items':sum(len(x['items']) for x in programs),'categories':len({x['category'] for x in programs}),'source_reference_files':len(evidence)},'programs':programs}
dump('review-data.json',data)
counts=collections.Counter(i['priority'] for p in programs for i in p['items'])
views=collections.Counter(i['perspective'] for p in programs for i in p['items'])

intro=f'''# AIMaster 프로그램별 상세 개선 분석 · 2026-10-19

**83개 프로그램 · 265개 세부 항목 · 14개 카테고리**. 각 항목은 필요 이유·현재 코드 조건·제안 범위·확인 시험을 갖습니다.

- [검색 보고서](REPORT.html) · [엑셀 표](AIMaster_프로그램별_상세개선_20261019.xlsx) · [공통 아이디어 원장](../PROGRAM_IDEA_BACKLOG.md) · [인수인계](../PROGRAM_DETAILED_REVIEW_HANDOFF_20261019.md) · [분석 중 최신 변경 재대조](LATEST_RECHECK.md)
- 기준 소스 `{BASE}`. 카탈로그 분류/주소는 앞선 분석 스냅샷을 재사용했습니다. 운영 DB 버전이나 배포 상태를 다시 인증한 표가 아닙니다.
- {data['scope']} 코드 조건 확인과 실제 장애 재현을 구분합니다. 운영 DB·회원 데이터·환경변수 변경·유료 호출·게시/발송은 **0회**입니다.
- 우선순위: **P0** 접근/통신 경계 우선 검토, **P1** 자료 격리·손실·중복 실행·비용 위험, **P2** 검증/이해/편의, **P3** 이후 적용할 기능 확장 후보. 우선순위는 실제 사고 발생 판정이나 작업 승인과 다릅니다.
- 세 관점은 개발·유지보수 / 디자인·사용성 / 신규 자동화입니다. 신규 자동화에는 기존 기능을 활용하는 작은 개선도 포함하며 세상에 없던 신규성을 주장하지 않습니다.
- 채택 기준 세 가지는 **안정적으로 동작하는가 / 현재 프로그램에 적용 가능한가 / 사용자가 편리한가**입니다. 근거 없는 점수 대신 항목별 검증 조건을 제시했습니다.
- 이미 해결된 기능은 보존하고 이전 문서와 충돌하면 최신 코드·완료 증거를 우선 대조합니다. 본인 계정/API·개인 자료 격리·생성물 최초 생성 +30일은 공통 경계입니다. 사용자 설정/키/권한/원본 파일은 생성 자료와 분류를 구분합니다.
- 당분간 기존 프로그램 안정화가 우선입니다. 검증된 기능만 기준 모듈 후보로 남기고 신규/리뉴얼에 한 개씩 적용합니다. 일괄 모듈화·앱 간 런타임 호출은 하지 않습니다.

## 프로그램별 찾아보기

'''
summary_rows=[]
for p in programs:
 summary_rows.append([p['sequence'],p['category'],f"[{p['name']}](programs/{p['slug']}.md)",len(p['items']),p['items'][0]['priority'],p['items'][0]['title'],p['items'][0]['why'],f"[프로그램]({p['app_url']})"])
write('INDEX.md',intro+table(['번호','카테고리','프로그램별 분석','항목','첫 검토 우선','첫 검토 주제','왜 필요한가','링크'],summary_rows)+'\n## 항목 분포\n\n'+table(['우선순위','항목 수'],sorted(counts.items()))+'\n'+table(['관점','항목 수'],views.items())+'\n## 사용 방법\n\n작업 전 공통 원장 → 이 목록 → 프로그램별 분석 → 최신 AGENTS/README·개선 원장·ACTIVE_WORK를 읽습니다. 같은 문제를 중복 구현하지 않고 최신 결과와 대조한 뒤 플랜을 제시합니다. 작업 후 프로그램 원장에 처리한 ID·버전·증거·미완료·다음 첫 단계를 추가하고 공통 원장에서 연결합니다. 이 분석서는 날짜가 고정된 근거 스냅샷이며 최신 완료 기록을 지우지 않습니다.\n')

for p in programs:
 text=f"# {p['name']} — 상세 개선 분석\n\n"
 text+=f"카테고리 **{p['category']}** · 폴더 `{p['folder']}` · [프로그램 열기]({p['app_url']}) · [전체 목록](../INDEX.md) · [공통 아이디어 원장](../../PROGRAM_IDEA_BACKLOG.md)\n\n"
 text+=f"**정적 검토 기준** `{BASE[:12]}` · 코드 버전 `{p['code_version_at_basis']}`. 운영 배포·DB 버전 인증이 아니며, 이후 변경은 최신 AGENTS/README가 우선입니다.\n\n"
 text+='## 현재 기능과 보존할 장점\n\n'+p['workflow']+'\n\n**보존할 장점:** '+p['strength']+'\n\n기능 설명은 기존 카탈로그·분석 기록을 재사용한 개요입니다. 이번 근거는 아래 선택 소스 경로의 조건이며 전체 메뉴/실제 기능을 모두 다시 실행한 것은 아닙니다.\n\n'
 text+='## 개선 항목 표\n\n'+table(['ID','우선순위','관점','개선 주제','왜 필요한가'],[[i['id'],i['priority'],i['perspective'],i['title'],i['why']]for i in p['items']])
 for i in p['items']:
  text+='\n## '+i['id']+' · '+i['title']+'\n\n'
  text+=table(['구분','상세 내용'],[
   ['우선순위·성격',i['priority']+' · '+i['perspective']+' · '+i['kind']],
   ['근거',f"[{i['evidence_file']}:{i['evidence_line']}]({i['evidence_url']}) · 기준 커밋 고정 링크"],
   ['현재 확인한 조건',i['current']],['확인 수준·미검증',i['confirmation']],['왜 필요한가',i['why']],
   ['구체적인 개선 범위',i['proposal']],['완료 판단·시험 표본',i['acceptance']],
   ['기준 1: 안정적 자동화',i['stability']],['기준 2: 현재 적용성',i['applicability']],['기준 3: 사용자 편의',i['convenience']],
   ['안전·비용·승인 경계',i['implementation_boundary']+' '+i['approval_note']],['현재 상태',i['status']]])
 text+='\n## 다음 CLI의 첫 단계\n\n1. 최신 공통 지침·ACTIVE_WORK·이 프로그램 AGENTS/README·개선 원장을 읽고 다른 CLI 변경과 대조합니다.\n2. 위 첫 항목부터 실제 남은 문제인지 확인하고, 이미 완료됐으면 완료 증거를 연결합니다. 수정 전 개발 플랜을 제시합니다.\n3. 모의 응답·합성 파일 등 비용 없는 재현을 우선합니다. 운영 권한/DB/외부 발송이 필요한 검사는 기존 승인 범위를 확인합니다.\n4. 기존 기능을 보존하며 한 프로그램씩 적용하고 실제 수정 후 빌드·검수·커밋·푸시·배포·운영 확인을 진행합니다. 문서 분석만 한 이번 작업에는 앱 배포가 없습니다.\n5. 처리한 ID·실제 시험·미검수·새 인사이트·라이브 링크를 프로그램 원장과 공통 인수인계에 남깁니다.\n'
 write('programs/'+p['slug']+'.md',text)

wb=Workbook();wb.remove(wb.active)
def sheet(name,headers,rows,widths=None):
 ws=wb.create_sheet(name);ws.append(headers)
 for row in rows:ws.append([str(v) if isinstance(v,(dict,list)) else v for v in row])
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='173B56');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='top')
 for row in ws.iter_rows(min_row=2):
  for c in row:
   if isinstance(c.value,str) and c.value.startswith(('=','+','-','@')):c.data_type='s'
   c.alignment=Alignment(wrap_text=True,vertical='top')
   if c.row%2==0:c.fill=PatternFill('solid',fgColor='F0F6FA')
  ws.row_dimensions[row[0].row].height=90
 for j in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(j)].width=(widths or {}).get(j,36 if j>3 else 22)
 return ws
sheet('읽는 방법',['항목','내용'],[['범위',data['scope']],['기준 소스',BASE],['프로그램/개선',f'83개 / 265개 / 14개 카테고리'],['구현 상태','제안과 검증 과제. 최신 프로그램 원장 우선.'],['비용/운영 변경','실제 유료 API·게시/발송·DB/키/환경변수 변경 0회'],['채택 기준','안정적인 자동화 / 현재 적용성 / 사용자 편의'],['적용 순서','현재 기능 안정화 → 검증된 기능 기준 후보 → 신규/리뉴얼에 하나씩 적용'],['협업','작업 전 원장 확인·최신 코드 대조·플랜; 작업 후 결과/미완료/인사이트/링크 기록'] ],{1:24,2:105})
ws=sheet('프로그램83',['번호','카테고리','프로그램','폴더','항목 수','코드 버전(기준)','첫 검토 주제','필요 이유','프로그램 링크','상세 문서'],[[p['sequence'],p['category'],p['name'],p['folder'],len(p['items']),p['code_version_at_basis'],p['items'][0]['title'],p['items'][0]['why'],p['app_url'],'programs/'+p['slug']+'.md']for p in programs],{7:48,8:80})
for row in ws.iter_rows(min_row=2):row[8].hyperlink=row[8].value;row[9].hyperlink=row[9].value
allitems=[(p,i)for p in programs for i in p['items']]
sheet('상세개선265',['ID','프로그램','우선순위','관점','개선 주제','현재 조건','필요한 이유','구체적 제안 범위','완료 기준·시험','확인 수준','근거 파일','행','근거 링크','상태'],[[i['id'],p['name'],i['priority'],i['perspective'],i['title'],i['current'],i['why'],i['proposal'],i['acceptance'],i['confirmation'],i['evidence_file'],i['evidence_line'],i['evidence_url'],i['status']]for p,i in allitems],{5:48,6:65,7:75,8:80,9:80})
sheet('채택기준',['ID','프로그램','안정적 자동화','현재 적용성','사용자 편의','안전·승인 경계'],[[i['id'],p['name'],i['stability'],i['applicability'],i['convenience'],i['approval_note']]for p,i in allitems],{3:85,4:65,5:75,6:75})
sheet('소스근거',['파일','기준 커밋','SHA256(UTF8)','전체 행 수'],[[v['path'],v['revision'],v['sha256_utf8'],v['lines']]for v in evidence.values()],{1:85,2:48,3:70})
logs=[json.loads(s)for s in (OUT/'read-log.jsonl').read_text(encoding='utf-8').splitlines()if s.strip()]
unique={json.dumps(x,sort_keys=True):x for x in logs}
sheet('선택열람기록',['파일','표시한 행(전체 정독 아님)','기준 커밋'],[[x['file'],str(x.get('displayed_lines')or[str(x.get('start'))+'~'+str(x.get('end'))]),x['revision']]for x in unique.values()],{1:85,2:65,3:48})
sheet('후속운영',['항목','완료/미완료·이유'],[['완료','83개 개별 분석·265개 구체 항목·이유·근거·시험 조건'],['미실행','265개 개선 구현·전 앱 빌드/운영 검수·DB 변경·유료 호출'],['최신성','문서 변경은 source-drift.json. 구현 전 최신 원장 확인'],['폴더 배포','문서 설치/보류는 publication-manifest.json 참조'],['공통 원장','docs/PROGRAM_IDEA_BACKLOG.md에서 이 보고서 연결'],['마켓플레이스 요청','현재 Codex 계정에 mattpocock/skills 소스 추가 확인. 개별 스킬은 설치하지 않음']],{1:26,2:110})
wb.save(OUT/'AIMaster_프로그램별_상세개선_20261019.xlsx')

# Self-contained report: no remote resources, safe text rendering for all data.
cards=[]
for p in programs:
 its=[]
 for i in p['items']:
  rows=[('현재 조건',i['current']),('왜 필요한가',i['why']),('제안 범위',i['proposal']),('완료 기준·시험',i['acceptance']),('확인 수준',i['confirmation']),('채택 기준',i['stability']+' / '+i['applicability']+' / '+i['convenience'])]
  its.append(f'<section class="item" data-priority="{i["priority"]}" data-view="{esc(i["perspective"])}"><h3><span class="badge {i["priority"]}">{i["priority"]}</span> {i["id"]} · {esc(i["title"])}</h3><p class="muted">{esc(i["perspective"])} · {esc(i["kind"])}</p><dl>'+''.join('<dt>'+esc(k)+'</dt><dd>'+esc(v)+'</dd>'for k,v in rows)+f'</dl><a href="{esc(i["evidence_url"])}" target="_blank" rel="noopener">근거 코드 · {esc(i["evidence_file"])}:{i["evidence_line"]}</a></section>')
 cards.append(f'<article class="program" data-category="{esc(p["category"])}" data-slug="{p["slug"]}"><details><summary><span class="num">{p["sequence"]:02}</span><span><b>{esc(p["name"])}</b><small>{esc(p["category"])} · {len(p["items"])}개 항목 · {esc(p["slug"])}</small></span></summary><div class="program-body"><p>{esc(p["short_desc"])}</p><p><strong>보존할 장점:</strong> {esc(p["strength"])}</p><p><a href="{esc(p["app_url"])}" target="_blank" rel="noopener">프로그램 열기 ↗</a> · <a href="programs/{p["slug"]}.md">상세 문서</a></p>'+''.join(its)+'</div></details></article>')
cats=''.join(f'<option>{esc(c)}</option>'for c in sorted({p['category']for p in programs}))
page='''<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>AIMaster 프로그램별 상세 개선 분석 · 2026-10-19</title><style>
:root{font-family:Segoe UI,Malgun Gothic,sans-serif;color:#183445;background:#f4f7fa;line-height:1.65}*{box-sizing:border-box}body{margin:0}main{max-width:1180px;margin:auto;padding:32px 24px}a{color:#075f9c;overflow-wrap:anywhere}header{background:#173b56;color:white;padding:32px;border-radius:20px}header a{color:#d5efff}h1{font-size:30px;line-height:1.4;margin:0 0 12px}h2{font-size:22px}h3{font-size:17px;margin:0}.lead{font-size:18px}.stats{display:grid;grid-template-columns:repeat(4,1fr);gap:12px;margin:20px 0}.stat{background:white;border:1px solid #d3e0e9;padding:20px;border-radius:12px}.stat b{display:block;font-size:30px;color:#096b82}.notice{padding:20px;background:#fff5df;border:1px solid #e4c58c;border-radius:12px}.muted,small{color:#4d6576}.controls{position:sticky;top:0;z-index:2;background:#f4f7faf5;padding:16px 0;display:grid;grid-template-columns:2fr 1fr 1fr 1.3fr;gap:10px}label{font-size:13px;font-weight:600}input,select,button{font:inherit;border:1px solid #a7bdcc;border-radius:8px;padding:10px;background:white;color:#183445;width:100%}button{width:auto;cursor:pointer}.actions{display:flex;gap:12px;align-items:center;margin:8px 0 18px;flex-wrap:wrap}.program{background:white;border:1px solid #ccdce7;margin:12px 0;border-radius:12px;overflow:hidden}summary{display:flex;gap:16px;padding:20px;cursor:pointer;align-items:center;list-style:none}summary:hover{background:#edf7fc}summary small{display:block}.num{background:#e5f0f6;padding:8px;border-radius:8px;font-size:18px}.program-body{padding:0 24px 24px}.item{padding:22px 0;border-top:1px solid #d9e3eb}.badge{display:inline-block;padding:2px 8px;font-size:12px;border-radius:6px;background:#e8eff5}.P0{background:#ffe0e0;color:#8e1717}.P1{background:#ffebce;color:#7a4500}.P2{background:#e0efff;color:#174a7b}.P3{background:#e1f2e9;color:#246145}dl{display:grid;grid-template-columns:120px minmax(0,1fr);gap:10px;margin:16px 0}dt{font-weight:700}dd{margin:0;overflow-wrap:anywhere}[hidden]{display:none!important}footer{margin-top:32px;font-size:13px;color:#4d6576}@media(max-width:700px){main{padding:16px 12px}header{padding:22px 18px}h1{font-size:23px}.stats{grid-template-columns:1fr 1fr}.stat{padding:12px}.controls{position:static;grid-template-columns:1fr 1fr}.controls label:first-child{grid-column:1/-1}summary{padding:16px;gap:12px}.program-body{padding:0 16px 18px}dl{grid-template-columns:1fr;gap:4px}dd{margin-bottom:12px}}
</style><main><header><p>2026.10.19 · CLI 공동 개선 검토 자료</p><h1>AIMaster 프로그램별 상세 개선 분석</h1><p class="lead">어디를 바꿀지, 왜 필요한지, 어떻게 확인할지</p><p>'''+f'{len(programs)}개 프로그램의 선택 핵심 소스·문서를 비교한 정적 분석입니다. 실제 운영 장애·전 프로그램 검수 완료를 뜻하지 않습니다.</p><a href="AIMaster_프로그램별_상세개선_20261019.xlsx">엑셀 내려받기</a> · <a href="INDEX.md">전체 목록</a> · <a href="../PROGRAM_IDEA_BACKLOG.md">공통 아이디어 원장</a></header><div class="stats"><div class="stat"><b>83</b>프로그램</div><div class="stat"><b>265</b>상세 항목</div><div class="stat"><b>14</b>카테고리</div><div class="stat"><b>3</b>검토 관점</div></div>'+'''
<section class="notice"><strong>지금은 기존 기능 안정화가 우선입니다.</strong><p>각 항목은 코드 조건·필요 이유·수정 범위·시험 조건을 구분합니다. 이미 해결된 기능은 보존하며 최신 원장과 대조합니다. 검증된 기능의 모듈화는 신규·리뉴얼에 하나씩 적용합니다.</p><p>채택 기준: 안정적인 자동화 / 현재 프로그램 적용성 / 사용자 편의. 본인 계정·API와 개인 자료 격리·생성물 최초 생성 +30일을 지킵니다. 제안은 구현·DB 변경 승인이 아닙니다.</p></section>
<div class="controls"><label>프로그램·개선 내용 검색<input id="search" type="search" placeholder="예: 예약, 보관, 부분 실패, 프로그램명"></label><label>카테고리<select id="category"><option value="">전체</option>'''+cats+'''</select></label><label>우선순위<select id="priority"><option value="">전체</option><option>P0</option><option>P1</option><option>P2</option><option>P3</option></select></label><label>검토 관점<select id="view"><option value="">전체</option><option>개발·유지보수</option><option>디자인·사용성</option><option>신규 자동화</option></select></label></div><div class="actions"><button id="reset">필터 초기화</button><button id="expand">표시된 프로그램 펼치기</button><output id="count" aria-live="polite"></output></div><div id="programs">'''+''.join(cards)+'''</div><p id="empty" hidden>조건에 맞는 항목이 없습니다. 검색어나 필터를 바꿔 주세요.</p><footer>분석 문서만 작성했습니다. 실제 유료 API·게시/발송·운영 DB/환경변수 변경 0회. 구현과 실제 회원 기능 검수는 후속 작업입니다.<br>'''+f'근거 커밋: {BASE}. 최신 변경 비교: {HEAD}. <a href="../PROGRAM_DETAILED_REVIEW_HANDOFF_20261019.md">완료·미완료 인수인계</a></footer></main>'+'''
<script>
const programs=[...document.querySelectorAll('.program')],s=document.querySelector('#search'),c=document.querySelector('#category'),p=document.querySelector('#priority'),v=document.querySelector('#view');
function filter(){const q=s.value.trim().toLocaleLowerCase();let pc=0,ic=0;for(const app of programs){let n=0;const appMatch=!c.value||app.dataset.category===c.value;const identity=app.querySelector('summary').textContent.toLocaleLowerCase();for(const item of app.querySelectorAll('.item')){const match=appMatch&&(!p.value||item.dataset.priority===p.value)&&(!v.value||item.dataset.view===v.value)&&(!q||(identity+' '+item.textContent.toLocaleLowerCase()).includes(q));item.hidden=!match;if(match)n++}app.hidden=n===0;if(n){pc++;ic+=n}}document.querySelector('#count').textContent=pc+'개 프로그램 · '+ic+'개 항목';document.querySelector('#empty').hidden=pc>0}
for(const el of [s,c,p,v])el.addEventListener('input',filter);document.querySelector('#reset').onclick=()=>{s.value=c.value=p.value=v.value='';filter()};document.querySelector('#expand').onclick=()=>{for(const app of programs)if(!app.hidden)app.querySelector('details').open=true};filter();
</script></html>'''
for key,label in [('category','카테고리'),('priority','우선순위'),('view','검토 관점')]:
 page=page.replace(f'<select id="{key}">',f'<select id="{key}" aria-label="{label}">')
write('REPORT.html',page)
print(json.dumps({'programs':len(programs),'items':len(allitems),'priority':dict(counts),'views':dict(views),'source_refs':len(evidence)},ensure_ascii=False))
