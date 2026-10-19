"""Meaningful artifact consistency checks; never runs product APIs or changes member data."""
import pathlib,json,re,csv,hashlib
from openpyxl import load_workbook
OUT=pathlib.Path(__file__).parent
data=json.loads((OUT/'proposal-data.json').read_text(encoding='utf-8'))
checks=[]
def check(name,condition):
 if not condition:raise AssertionError(name)
 checks.append({'name':name,'passed':True})
check('83 unique program slugs',len(data['programs'])==83 and len({x['slug'] for x in data['programs']})==83)
check('14 categories and 81/2 activation',len({x['category'] for x in data['programs']})==14 and sum(x['is_active'] for x in data['programs'])==81)
check('83 program-specific recommendations',all(x['idea'] and x['strength'] and x['review_task'] for x in data['programs']))
check('19 unique issue IDs / 3 P0',len(data['issues'])==19 and len({x['id'] for x in data['issues']})==19 and sum(x['priority']=='P0' for x in data['issues'])==3)
check('12 merge proposals / 18 creative ideas / 12 work units',[len(data[x]) for x in ['merges','ideas','roadmap']]==[12,18,12])
known={x['slug'] for x in data['programs']}
check('Merge source programs all registered',all(set(x['sources'])<=known for x in data['merges']))
paths=[]
for x in data['issues']:
 for ref in x['refs']:
  p=OUT/re.sub(r':\d+$','',ref);paths.append((ref,p.exists()))
check('All finding reference paths exist',all(ok for ref,ok in paths))
check('83 program entry documents exist',all((OUT.parents[1]/p['doc_entry']).exists() for p in data['programs']))
with (OUT/'programs.csv').open(encoding='utf-8-sig',newline='') as f:rows=list(csv.reader(f))
check('CSV exact 83 rows / 17 columns / BOM',len(rows)==84 and all(len(r)==17 for r in rows) and (OUT/'programs.csv').read_bytes().startswith(b'\xef\xbb\xbf'))
wb=load_workbook(OUT/'AIMaster_분석_개발제안_20261019.xlsx',read_only=False,data_only=False)
expected={'프로그램83':84,'問題点19':20,'병합제안12':13,'아이디어18':19,'협업로드맵12':13}
expected['문제점19']=expected.pop('問題点19');expected['레퍼런스모듈18']=19
check('Excel 8 sheets and exact table row counts',len(wb.sheetnames)==8 and all(wb[k].max_row==n for k,n in expected.items()))
check('Excel frozen/filter headers',all(w.freeze_panes=='A2' and w.auto_filter.ref for w in wb.worksheets))
check('Excel no executable formulas',not any(c.data_type=='f' for w in wb for row in w for c in row))
check('Excel exact 83 live hyperlinks',sum(bool(c.hyperlink) for row in wb['프로그램83'] for c in row)==83)
metadata=json.loads((OUT/'db-detail-evidence.json').read_text(encoding='utf-8'))['data'][0]['jsonb_build_object']
check('F001 evidence actually has four PUBLIC ALL grants',len(metadata['critical_grants'])==4 and all(x['roles']==['public'] and x['command']=='ALL' and x['using']=='true' and x['check']=='true' and all(x[k] for k in ['anon_select','anon_insert','anon_update','anon_delete']) for x in metadata['critical_grants']))
check('F007 response-id unique index absent in inspected metadata',not any('UNIQUE' in x['definition'] and 'response_id' in x['definition'] for x in metadata['crm_indexes']))
tool=json.loads((OUT/'tooling-evidence.json').read_text(encoding='utf-8'))
check('Tool findings actually reproduced',tool['lint']['exit']==1 and tool['booking_timezone']['different'] and tool['shared_module']['manual_timeout_checks']==6)
http=json.loads((OUT/'http-evidence.json').read_text(encoding='utf-8'))
check('83 HTTP records and only empty-input gate checks',len(http['availability'])==83 and len(http['empty_input_gates'])==2 and all(x['status']==400 for x in http['empty_input_gates']))
browser=json.loads((OUT/'report-browser-verification.json').read_text(encoding='utf-8'))
check('Actual headed Edge report checks passed',len(browser['checks'])==14 and all(x['passed'] for x in browser['checks']) and not browser['errors'] and not browser['external_network'])
modules=data['reference_modules']
check('18 reference modules distinguish 4 registered / 14 candidates',len(modules)==18 and sum(m['status']=='기존 등록' for m in modules)==4)
check('Every module reference path exists',all((OUT.parents[1]/p.strip()).exists() for m in modules for p in m['references'].split(';')))
ideas=(OUT.parent/'PROGRAM_IDEA_BACKLOG.md').read_text(encoding='utf-8')
check('Ongoing ideas preserve 18 proposal IDs plus user direction',len(re.findall(r'^\| (?:I-\d\d|U-01) \|',ideas,re.M))==19 and all(k in ideas for k in ['안정적인 자동화','기존 프로그램 적용성','사용자 편의']))
begin='<!-- INCREMENTAL_REFERENCE_IDEAS_20261019 -->';end='<!-- /INCREMENTAL_REFERENCE_IDEAS_20261019 -->'
blocks=[]
for name in ['AGENTS.md','CLAUDE.md','GEMINI.md']:
 t=(OUT.parents[1]/name).read_text(encoding='utf-8');blocks.append(t[t.index(begin):t.index(end)+len(end)])
check('Three CLI entry rules carry identical incremental adoption / idea policy',len(set(blocks))==1 and '일괄 수정' in blocks[0])
for filename in ['REPORT.html','REPORT.md','ROADMAP.md','README.md','REFERENCE_MODULE_PROPOSAL.md']:
 txt=(OUT/filename).read_text(encoding='utf-8');check(filename+' Korean encoding/content',bool(re.search('[가-힣]',txt)) and '\ufffd' not in txt)
result={'checks':checks,'passed':len(checks),'scope':'Report content/metadata/Excel/CSV/local HTML only, not product functionality certification','artifact_sha256':{p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [OUT/'REPORT.html',OUT/'REPORT.md',OUT/'AIMaster_분석_개발제안_20261019.xlsx',OUT/'programs.csv']}}
(OUT/'artifact-verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({'passed':len(checks),'paths':len(paths)},ensure_ascii=False))
