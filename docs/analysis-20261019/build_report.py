import json,pathlib,re,collections,csv,html,datetime,subprocess,hashlib
from proposal_data import PROGRAM_NOTES,ISSUES,MERGES,IDEAS,ROADMAP
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo

OUT=pathlib.Path(__file__).parent;ROOT=OUT.parents[1]
inventory=json.loads((OUT/'inventory.json').read_text(encoding='utf-8')); programs=inventory['programs']; base=inventory['source_revision']
updates=json.loads((OUT/'updates-during-review.json').read_text(encoding='utf-8')) if (OUT/'updates-during-review.json').exists() else {'changes':[]}
for p in programs:
 change=next((u for u in updates['changes'] if u['slug']==p['slug']),None)
 if change:
  p['version']=change['final_db']
  for v in p['code_versions']:v['version']=change['final_code']
notes={v[0]:v[1:] for line in PROGRAM_NOTES.strip().splitlines() if (v:=line.split('|'))}
assert len(notes)==83 and set(notes)=={p['slug'] for p in programs},'Missing program-specific analyst notes'
db=json.loads((OUT/'db-metadata.json').read_text(encoding='utf-8'))['data'][0]['jsonb_build_object']
http=json.loads((OUT/'http-evidence.json').read_text(encoding='utf-8'))
family={'데이터수집|기획':'리서치·문서·기획','쓰레드':'SNS 콘텐츠·소통','인스타':'SNS 콘텐츠·소통','블로그':'블로그·네이버 운영','네이버':'블로그·네이버 운영','유튜브':'유튜브 성장·제작','CRM':'고객관계·마케팅','마케팅(홍보)':'고객관계·마케팅','이커머스':'상품소싱·커머스','이미지':'이미지·브랜드 디자인','음악':'음악·미디어 제작','멀티미디어':'음악·미디어 제작','부동산':'부동산 데이터 분석','Tools':'문서·업무 생산성'}
for p in programs:
 n=notes[p['slug']];p.update(strength=n[0],review_task=n[1],idea=n[2],integration=n[3],proposed_family=family[p['category']],code_version=', '.join(v['version'] for v in p['code_versions']) or '미발견')
 if p['slug'] in ['personality-suite','personality-code','mbti-character','tarot-reading']:p['proposed_family']='라이프·자기탐색'
 p['issue_ids']=[x['id'] for x in ISSUES if p['slug'] in x['targets']]
 p['merge_ids']=[x['id'] for x in MERGES if p['slug'] in x['sources']]
 p['doc_entry']=p['folder']+'/AGENTS.md' if (ROOT/p['folder']/'AGENTS.md').exists() else p['folder']+'/README.md'
 p['feature']=p['baseline_capabilities'].get('기능') or p['short_desc'] or p['strength']
 p['coverage']='카탈로그·문서·소스 정적 점검; 핵심 경로 추가 검토' if p['issue_ids'] else '카탈로그·문서·소스 정적 점검; 개별 실행 미검수'
counts=collections.Counter(p['category'] for p in programs);families=collections.Counter(p['proposed_family'] for p in programs)
evidence={'generated_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'source_revision':base,'scope':'83 catalog programs; fixed committed-source inventory + selected manual code review and metadata queries. Not all programs fully executed.','programs':83,'active':sum(p['is_active'] for p in programs),'inactive':sum(not p['is_active'] for p in programs),'categories':dict(counts),'proposed_families':dict(families),'source_files_screened':sum(p['source_files'] for p in programs),'source_lines_screened':sum(p['source_lines'] for p in programs),'documents_indexed':sum(len(p['docs']) for p in programs),'test_files_found':sum(p['test_files'] for p in programs),'programs_without_test_folder_match':sum(p['test_files']==0 for p in programs),'issues':len(ISSUES),'merges':len(MERGES),'ideas':len(IDEAS),'roadmap':len(ROADMAP),'http_200_after_redirect':sum(v['status']==200 for v in http['availability']),'rls_tables':sum(v['rls'] for v in db['tables']),'db_tables':len(db['tables']),'pg_cron_jobs_observed':len(db['cron']),'paid_program_api_calls':0,'external_publish_send':0,'member_rows_read':0,'product_code_changes':0,'db_mutations':0,'deployment_builds':0}
(OUT/'summary.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
data={'summary':evidence,'programs':[{k:v for k,v in p.items() if k not in ['signals','routes','packages','baseline_capabilities']} for p in programs],'issues':ISSUES,'merges':MERGES,'ideas':[dict(zip(['id','name','scope','description','guardrail','acceptance'],v)) for v in IDEAS],'roadmap':[dict(zip(['id','priority','work','issues','dependencies','roles','approval','acceptance'],v)) for v in ROADMAP]}
(OUT/'proposal-data.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')

def md(x):return str(x if x is not None else '').replace('|',' / ').replace('\n',' ')
def table(headers,rows):return '\n'.join(['| '+' | '.join(map(md,headers))+' |','|'+'|'.join(['---']*len(headers))+'|',*['| '+' | '.join(map(md,r))+' |' for r in rows]])+'\n'
def src(ref):
 path=re.sub(r':\d+$','',ref);return f'[{ref}]({path})'

intro=f'''# AIMaster 프로그램 정밀 분석·추가 개발 제안서

2026-10-19 · Codex · **분석 및 제안만 작성 / 운영 수정 미실행**

**판단:** 신규 기능을 늘리기 전에 공통 DB 접근 정책·직접 API 보호·회원별 보관·발송 정합성을 먼저 정비하고, 검증된 엔진을 활용한 작은 연계 MVP부터 확장하는 순서를 권합니다. 기존 상품을 폐지하거나 회원 간 자료를 공유하는 병합은 제안하지 않습니다.

[보기 좋은 HTML 보고서](REPORT.html) · [엑셀 전체 표](AIMaster_분석_개발제안_20261019.xlsx) · [프로그램 CSV](programs.csv) · [협업 실행표](ROADMAP.md)

## 분석 범위·확인 수준

- DB 카탈로그 **83개 / 활성81·비활성2 / 현재14개 카테고리**. 기능 중심 보조분류11개를 별도로 제안합니다. DB 카테고리를 수정한 것은 아닙니다.
- 고정 소스 기준 **{base[:8]}**, {evidence['source_files_screened']:,}개 소스 파일·{evidence['source_lines_screened']:,}행을 자동 수집/정적 패턴 점검하고 프로그램 문서 {evidence['documents_indexed']}개를 색인했습니다. 전 코드를 사람이 한 줄씩 심층 검토하거나 전 기능을 실행했다는 뜻이 아닙니다. 위험 신호가 있는 주요 경로는 파일/호출/DB정책을 직접 대조했습니다.
- 운영 DB는 **테이블202개 모두 RLS 켜짐**을 확인했지만 RLS 이름/활성 여부만으로 소유자 격리를 인정하지 않았습니다. 정책·역할·grant를 함께 확인했습니다. 회원 행·API키/비밀번호 원문은 읽지 않았습니다.
- 등록 주소83개는 비로그인 HEAD·리디렉트 후200. 로그인 화면도200이므로 기능 정상 인증이 아닙니다. eBook API는 비용/조회가 없는 빈 입력2개만 POST해 인증보다 입력검증이 먼저 실행됨을 확인했습니다.
- 이번 실제 실행은 읽기 전용 메타데이터/공용 사본 검사/무해한 HTTP/오디오 lint 재현입니다. 전 프로그램 빌드·npm audit·회원 로그인 E2E·실제 수집/AI/게시/발송·렌더는 미실행. 제품 API 유료 호출0·DB변경0·운영배포0입니다.
- **확인 수준:** 운영 설정 확인 / 코드 확인(논리적 경로) / 조건부 코드 결함 / 기존 인수인계 대조 / 신규 제안. '버그가 없다'는 인증이나 실제 피해 발생을 뜻하지 않습니다.
- 다른 CLI가 계속 개발 중이므로 버전은 수집 시점 스냅샷입니다. 최종 재조회로 ai-task-planner v1.08·banana-canvas v1.10·landing-architect v1.10은 표에 갱신했습니다. 정적 점검량은 처음 고정 소스 기준입니다. [분석 중 갱신 기록](updates-during-review.json). 착수 전 최신 HEAD·DB·ACTIVE_WORK와 다시 비교합니다.

## 최우선 결론

{table(['순서','조치','대상','근거','승인·완료 기준'],[
('1','PUBLIC 전체 허용 DB 정책 정비','메인/모든 회원','F-001 운영 정책/anon CRUD 확인','DB 승인 → 역할별 allow/deny → 롤백/운영 검증'),
('2','직접 API 자격과 URL 안전 수집','eBook Studio','F-002/F-003','401/403·가짜 DNS/redirect·본문 상한 회귀'),
('3','회원 생성물 격리와 30일 보관','쇼핑제휴·쇼츠·포토·경쟁사','F-004~006/F-016','소유자·원생성일·직접조회 cutoff·비접속 정리'),
('4','유료 발송 중복/누락·회원 자격','CRM·예약·메일','F-007~011','선점·채널별 상태·정지/만료/비밀누락 차단'),
('5','새 프로그램은 작은 MVP부터','브랜드·자료정제·UI명세·사진규격','M-08/M-10/M-11/M-12','핵심2~3기능·기존 상품 유지·비용/권한 검증')])}

## 현재 카테고리별 분류

현재 분류를 보존한 표입니다. 추천 분류는 작업 경로 탐색용 보조 태그이며 상품 이동/삭제 지시가 아닙니다.

{table(['현재 카테고리','수','활성/비활성','보조 분야(안)','분류 개선 이유'],[(c,n,f"{sum(p['is_active'] for p in programs if p['category']==c)}/{sum(not p['is_active'] for p in programs if p['category']==c)}",family[c]+' (Tools의 자기점검4개 별도)', '목적·산출물·필요한 계정 태그를 함께 제공' if c not in ['데이터수집|기획','Tools'] else '수집/기획/UI제작 또는 문서/자기점검이 섞여 있어 보조 태그 필요') for c,n in counts.items()])}

보조 분야 집계: {', '.join(f'{c} {n}개' for c,n in families.items())}. 각 프로그램은 하나의 기본 보조분야와 여러 연계 태그를 가질 수 있습니다.

## 프로그램별 분석 표 — 83개 전수

강점은 현재 구조와 기존 기능 자료를 종합한 요약입니다. '검증 과제'에는 구조적 한계/아이디어도 포함하며, 확인된 결함은 F-ID를 붙였습니다. F-ID가 없다고 결함이 없다는 뜻은 아닙니다. 코드 버전/DB 버전은 독립 표시합니다.

'''
parts=[intro]
for c in counts:
 parts+=['### '+c+'\n\n',table(['프로그램·운영 링크','DB / 코드','강점','검증 과제·개선 기회','개선 아이디어','연계·병합 후보','근거·검토 수준'],[(f"[{p['name']}]({p['app_url']})"+(' **비활성**' if not p['is_active'] else ''),p['version']+' / '+p['code_version'],p['strength'],p['review_task'],p['idea'],p['integration']+'; '+','.join(p['merge_ids']),','.join(p['issue_ids'])+'; '+f"[프로그램 문서](../../{p['doc_entry']}); 정적점검") for p in programs if p['category']==c]),'\n']
parts+=['## 확인된 문제·조건부 위험·후속 조사 — '+str(len(ISSUES))+'건\n\n',table(['ID/우선','문제','대상','확인 수준','권장 방향'],[(x['id']+' '+x['priority'],x['title'],','.join(x['targets']),x['status'],x['fix']) for x in ISSUES]),'\n']
for x in ISSUES:
 parts += [f"### {x['id']} · {x['priority']} · {x['title']}\n\n",table(['구분','내용'],[('대상/수준',','.join(x['targets'])+' / '+x['status']),('관찰',x['observation']),('영향·한계',x['impact']),('개선',x['fix']),('검증 완료 기준',x['verify']),('승인·현재 상태',x['approval']+' / 제안서 기록만, 미수정'),('근거',' · '.join(src(r) for r in x['refs']))]),'\n']
parts+=['## 장점을 모은 병합 프로그램 후보 — 12개\n\n기존 상품은 유지합니다. 공용 엔진은 기준에서 검증한 사본을 새 독립 폴더에 가져오고, 기존 프로그램 API를 서로 호출하는 의존 구조를 만들지 않습니다. 병합은 장점만 공짜로 얻는 작업이 아니므로 추가 복잡도·비용·권한·회귀 범위를 함께 평가합니다. M/L/XL은 상대적 규모이며 일정/금액 견적이 아닙니다.\n\n',table(['ID·추천 단계','병합안','활용 프로그램','가져올 장점','줄일 단점','첫 MVP','선행조건·비용/위험','규모·통과 기준'],[(m['id']+' '+m['phase'],m['name'],', '.join(m['sources']),m['keep'],m['remove'],m['mvp'],m['risk'],m['effort']+' / '+m['gate']) for m in MERGES]),'\n',
'''## 연계 설계 원칙

1. **회원 본인만:** 공통 Supabase 한 개 안에서 user_id·RLS·개별 경로. 개인 키·외부 앱/계정은 수신 프로그램에서도 본인 것으로 확인합니다. 운영자 키/다른 회원 결과 폴백 없음.
2. **독립 배포:** 공통 엔진 기준→사본 검사→각 앱 독립 실행. 프로그램끼리 내부 API를 서로 호출해 한쪽 장애가 다른 쪽에 번지는 설계는 하지 않습니다.
3. **처음은 수동 전달:** 사용자가 JSON/ZIP을 명시적으로 내보내고 가져옵니다. schemaVersion/source/target/originalCreatedAt/expiresAt/checksum/rights 같은 명세를 두고, ownerHint는 인증으로 신뢰하지 않습니다. 키·토큰은 전달하지 않습니다. DB 기반 전달은 별도 승인 설계입니다.
4. **30일의 시작을 유지:** 이미 만든 결과를 다른 프로그램으로 복사/가져와 기간을 새로 시작하지 않습니다. 설정/개인 프롬프트/원본 파일/새 생성물의 분류를 먼저 확정합니다. 새로 생성된 파생물도 원자료 만료 이후 원자료를 재사용하지 않도록 계보/삭제 범위를 설계합니다.
5. **사람 확인:** 자료 이동만으로 AI 생성·게시·문자 발송을 시작하지 않습니다. 실제 외부 효과는 대상·비용·권리·수신 동의를 검토한 후 해당 회원의 명시적 실행으로 진행합니다.
6. **컨셉 보존:** 이미지 공용화는 모델·호출·대기·결과 수신까지. 프로그램별 이미지 프롬프트 내용/스타일은 유지합니다. 기존 UI/상품 메뉴/권한 제거는 별도 결정입니다.
7. **관리 가능한 범위:** 서버/앱 저장소의 조회차단·자동정리와 사용자가 다운로드한 개인 파일을 구분합니다. 브라우저를 켜지 않는 동안 IndexedDB 물리 삭제나 내려받은 파일 원격 삭제를 보장하지 않습니다.

## 창의적 확장 아이디어 — 18개

''',table(['ID','아이디어','적용 분야','사용자 가치·동작','설계 경계','검수 기준'],IDEAS),'\n',
'''## CLI 협업 실행 제안

아래는 역할 제안이며 현재 다른 CLI에 작업을 보내거나 담당을 확정하지 않았습니다. Codex는 기획·시스템 설계·오류 재현·디버깅·검수, Claude Code는 복잡한 구현·기능·리팩터링, Antigravity는 최신 지침의 UI/UX·매뉴얼 시각화·개선 기록을 중심으로 합니다. 실제 세부 담당은 주인님 배정/작업표로 조율합니다. 같은 프로그램 동시 작업 금지, 루트 공통 문서는 본인 항목만 추가합니다.

''',table(['작업ID','우선','작업','이슈/후보','선행','권장 역할','승인 조건','완료 기준'],ROADMAP),'\n',
'''## 다음 CLI가 착수할 때

- 최신 AGENTS/CLAUDE/GEMINI·TOP_RULE_PERSONAL_ACCOUNT_API·SHARED_FEATURE_MAP·프로그램 AGENTS/README·ACTIVE_WORK를 읽습니다. 작업 중이면 해당 폴더를 건드리지 않습니다.
- 같은 F-ID/W-ID에 실제 코드/운영 증거와 차이를 추가하고, 원문을 지우지 않습니다. 이미 해결됐다면 재현/완료 SHA로 '해소됨'을 기록합니다.
- 다른 CLI 결과와 차이가 나면 목적·최신 지시·검증 근거로 비교합니다. 기존 작업 대체/판단 어려움은 원본을 보존하고 주인님께 확인합니다.
- DB/환경변수/키·기존 자료 삭제·유료 대량 호출은 사전 승인. 제안서 기록은 구현 승인으로 해석하지 않습니다. 일반 기능도 개발 플랜을 먼저 보여드립니다.
- 검수는 본인 일반 시험 회원, Codex/Antigravity는 실제 Edge. 무료/합성 응답 우선. 실제 호출·미실행·비용을 구분하고 완료/미완료·라이브 링크·커밋/푸시/배포를 남깁니다.
- 배포는 최신 비용 절감 지침에 따라 필요한 프로그램만 묶어 하루1회, 급한 보안/이중과금은 예외 근거를 기록합니다. 이번 문서만 변경했으므로 앱 배포/DB버전은 필요 없습니다.

## 이미 해결됐거나 이번에 결함으로 보지 않은 항목

''',table(['항목','현재 판단','근거'],[
('부동산 상세 열람 자동 AI 비용','v1.13에서 명시적 실행으로 바뀜. 과거 원장의 미해결 문구를 새 결함으로 재등록하지 않음','real_estate_sales 상세 page.tsx:26'),
('YVS 전역 즐겨찾기·자동검색·85점 기본값','v1.14~후속 개선 적용. 이번 미해결은 다른 프로그램의 유사 경로','youtube-viral-studio 최신 AGENTS/Provider/보관DB'),
('YVS 1000개 제한·개발 high5','도달 시 안내 A안/경고 수용 판단 기록 있음. 새 결함/무조건 대규모 페이지 구현으로 바꾸지 않음','IMPROVEMENT_BACKLOG 최신 상태'),
('이미지 실패 안내·12p 이미지60초','eBook/웹툰은 실패 이유 전달이 추가됐고12p 이미지 route는300초. 웹툰 외부 대체 자체는 F-012로 분리','ebook cover·webtoon imageGen·12p generate-image route'),
('SNS 예약5분','현재 Threads/Instagram vercel.json이5분. 예전1일 예약 문제를 미해결로 중복 등록하지 않음','각 vercel.json'),
('file-merge 브라우저 검사 실패','최근 인수인계에서 서버 시작 순서 문제로12/12 확인. 제품 결함 아님','2026-10-19 HANDOFF·커밋4e574906'),
('사진 보관함 회원 격리','photo-remake는 UUID별IndexedDB 적용. 소유자 혼합으로 단정하지 않고 날짜 만료만 F-006','photo-remake history.ts'),
('네이버 운영자 키','trend action은 본인 키 조회. 전역캐시 문제와 분리','threads-affiliate-poster actions/trend.ts')]),'\n',
'''## 근거·자료·검수 한계

- [카탈로그 스냅샷](catalog-db.json), [정적 인벤토리](inventory.json), [DB 메타데이터](db-metadata.json), [DB 상세](db-detail-evidence.json), [HTTP 결과](http-evidence.json), [도구 재현](tooling-evidence.json), [공용 모듈 검사](shared-module-check.txt).
- [기존 역량지도](../PROGRAM_CAPABILITY_MAP.md), [공용 기능 지도](../SHARED_FEATURE_MAP.md), [공용 모듈](../SHARED_MODULES.md), [개선 기록 정책](../IMPROVEMENT_HANDOFF_POLICY.md), [DB 연동 정책](../DATABASE_CONNECTION_POLICY.md). 기존 역량지도 일부는2026-10-16의82개 기준이므로 기능 이해 보조로만 쓰고 현재 버전/활성/주소는 운영DB 스냅샷을 사용했습니다.
- Supabase의 RLS는 grant와 정책을 함께 평가하며 서비스 역할은 RLS를 우회합니다. 이번 DB 판단의 공식 기준: [RLS 문서](https://supabase.com/docs/guides/database/postgres/row-level-security). 정책 메타데이터로 확인한 허용 경로와 실제 데이터 피해는 구분했습니다.
- 자동정리는 실행 등록뿐 아니라 성공/실패 관측이 필요합니다. [Supabase Cron](https://supabase.com/docs/guides/cron). 조회된 pg_cron2개 외 Vercel Cron/외부 스케줄이 있을 수 있어 전체 정리 커버리지2개라고 단정하지 않았습니다.
- YouTube 할당량은 실제 회원 Cloud 프로젝트를 기준으로 표시합니다. [공식 할당량 문서](https://developers.google.com/youtube/data/determine_quota_cost). 창당 검색 호출/다른 API 단위를 구분하는 설계를 권하며 이번에 실제 YouTube를 호출하지 않았습니다.
- 모델/업로드/실행 시간은 제품 배포 설정과 제공사 한도를 함께 확인합니다. [Vercel Functions 한도](https://vercel.com/docs/functions/limitations), [WebCodecs](https://developer.mozilla.org/en-US/docs/Web/API/WebCodecs_API). 이 제안서가 모든 모델/브라우저 지원을 실측 보증하지 않습니다.
- 신규 프로그램/병합안은 구현 전 개별 플랜·기존 상품 권한·비용·데이터 정책을 확정해야 합니다. 현재 서버/회원 기능·키·DB·게시를 바꾸지 않았습니다.
''']
(OUT/'REPORT.md').write_text(''.join(parts),encoding='utf-8')
(OUT/'ROADMAP.md').write_text('# CLI 협업 실행표\n\n2026-10-19 제안·미배정·미구현. [전체 보고서](REPORT.md). 작업 전 최신 ACTIVE/승인 범위 확인.\n\n'+table(['ID','우선','작업','이슈','선행','권장 역할','승인','완료 기준'],ROADMAP)+'\n기존 결과 대체·판단 어려움은 원본을 보존하고 주인님께 확인합니다. 작업 결과에는 source/운영SHA·DB버전·실행/미실행 검수·비용·완료/미완료·서비스 링크를 기록합니다.\n',encoding='utf-8')

program_headers=['현재카테고리','보조분야안','프로그램','slug','활성','DB버전','코드버전','강점','검증과제','개선아이디어','연계','이슈ID','병합ID','운영URL','프로그램문서','발견테스트파일수','확인수준']
program_rows=[[p['category'],p['proposed_family'],p['name'],p['slug'],'활성' if p['is_active'] else '비활성',p['version'],p['code_version'],p['strength'],p['review_task'],p['idea'],p['integration'],','.join(p['issue_ids']),','.join(p['merge_ids']),p['app_url'],p['doc_entry'],p['test_files'],p['coverage']] for p in programs]
with (OUT/'programs.csv').open('w',encoding='utf-8-sig',newline='') as f:w=csv.writer(f);w.writerow(program_headers);w.writerows(program_rows)

# Offline, self-contained HTML. All DB/code values are escaped; no external script/font/CDN.
def h(x):return html.escape(str(x if x is not None else ''),quote=True)
def a(url,label):return f'<a href="{h(url)}" target="_blank" rel="noopener noreferrer">{h(label)}</a>'
def ht(headers,rows,attrs=None):
 body=[]
 for i,row in enumerate(rows):
  extra=attrs[i] if attrs else {}; extra['data-search']=' '.join(re.sub('<[^>]*>',' ',str(c)) for c in row)
  body.append('<tr '+' '.join(f'{h(k)}="{h(v)}"' for k,v in extra.items())+'>'+''.join('<td>'+str(c)+'</td>' for c in row)+'</tr>')
 return '<div class="table-wrap"><table><thead><tr>'+''.join('<th>'+h(t)+'</th>' for t in headers)+'</tr></thead><tbody>'+''.join(body)+'</tbody></table></div>'
cat_html=ht(['현재 분류','수','보조 분야(안)'],[[h(c),str(n),h(family[c])] for c,n in counts.items()])
prows=[];attrs=[]
for p in programs:
 prows.append([a(p['app_url'],p['name'])+f'<small>{h(p["slug"])} · '+('활성' if p['is_active'] else '비활성')+'</small>',h(p['category'])+'<small>'+h(p['proposed_family'])+'</small>',h(p['version'])+'<small>코드 '+h(p['code_version'])+'</small>',h(p['strength']),h(p['review_task']),h(p['idea']),h(p['integration'])+'<small>'+h(', '.join(p['merge_ids']))+'</small>',h(', '.join(p['issue_ids']) or '정적 기초 점검')+a('../../'+p['doc_entry'],'문서')])
 attrs.append({'data-category':p['category'],'data-state':'active' if p['is_active'] else 'inactive'})
program_html=ht(['프로그램/링크','분류','버전','강점','검증 과제/기회','개선 아이디어','연계/병합','근거'],prows,attrs)
issue_html=''
for x in ISSUES:
 issue_html+=f'<article id="{x["id"]}" class="issue {x["priority"]}" data-search="{h(x["id"]+x["title"]+str(x["targets"])+x["observation"])}"><h3><span class="badge {x["priority"]}">{x["priority"]}</span> {x["id"]} {h(x["title"])}</h3><p class="muted">{h(", ".join(x["targets"]))} · {h(x["status"])} · 이번 미수정</p>'+ht(['구분','내용'],[[h(k),h(v)] for k,v in [('관찰',x['observation']),('영향·제약',x['impact']),('개선 방향',x['fix']),('검수 기준',x['verify']),('승인',x['approval'])]])+'<p class="sources">'+' · '.join(a(re.sub(r':\d+$','',r),r) for r in x['refs'])+'</p></article>'
merge_html=ht(['후보','재사용 원천','가져올 장점/줄일 단점','첫 MVP','위험/선행','검수·규모'],[[f'<b>{m["id"]} {h(m["name"])}</b><small>{h(m["phase"])}</small>',h(', '.join(m['sources'])),h(m['keep'])+'<small>개선 목표: '+h(m['remove'])+'</small>',h(m['mvp']),h(m['risk']),h(m['gate'])+'<small>상대 규모 '+m['effort']+'</small>'] for m in MERGES])
idea_html=ht(['ID/아이디어','분야','동작/가치','설계 경계','완료 기준'],[[h(v[0]+' '+v[1]),h(v[2]),h(v[3]),h(v[4]),h(v[5])] for v in IDEAS])
road_html=ht(['작업/우선','내용','선행','권장 역할','승인 조건','완료 기준'],[[h(v[0]+' '+v[1]),h(v[2])+'<small>'+h(v[3])+'</small>',h(v[4]),h(v[5]),h(v[6]),h(v[7])] for v in ROADMAP])
critical_html=ht(['우선','조치','범위'],[['<b class="danger">P0</b>','DB PUBLIC 전체허용 / eBook API 자격·URL 보호','F-001~F-003'],['P1','회원격리·30일 보관 / 유료 발송 중복·누락 / 오디오·대체엔진','F-004~F-014·F-016'],['P2','버전·문서·검사 신뢰도와 작은 연계 MVP','F-015·F-017~F-019 / M 후보']])
html_doc='''<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>AIMaster 프로그램 분석·개발 제안서 · 2026-10-19</title><style>
:root{--ink:#182438;--muted:#58677e;--line:#dce4ef;--blue:#354fd8}*{box-sizing:border-box}body{margin:0;background:#f4f7fc;color:var(--ink);font:14px/1.7 "Malgun Gothic",sans-serif}header{background:#152743;color:white;padding:34px max(24px,calc((100vw - 1540px)/2))}header p{color:#c7d7ef;margin:10px 0}h1{font-size:30px;line-height:1.4;margin:0}h2{font-size:23px;margin:0 0 14px}h3{font-size:18px}a{color:#2f4ac7;text-underline-offset:3px}header a{color:#d5e5ff}main{max-width:1600px;margin:24px auto;padding:0 24px}.kpis{display:grid;grid-template-columns:repeat(6,1fr);gap:12px;margin:20px 0}.kpi{background:white;border:1px solid var(--line);border-radius:12px;padding:16px}.kpi b{display:block;font-size:27px;color:#263e98}.kpi span,.muted{color:var(--muted)}nav{display:flex;gap:7px;flex-wrap:wrap;position:sticky;top:0;background:#f4f7fc;padding:12px 0;z-index:5}button,input,select{font:inherit;border:1px solid #c9d4e6;background:white;color:var(--ink);border-radius:8px;padding:9px 12px}button{cursor:pointer}button[aria-selected=true]{background:var(--blue);color:white;border-color:var(--blue)}.panel{background:white;border:1px solid var(--line);border-radius:14px;padding:24px;margin-bottom:20px}.toolbar{display:flex;gap:10px;flex-wrap:wrap;margin-bottom:14px}.toolbar input{flex:1;min-width:200px}.table-wrap{overflow:auto;border:1px solid var(--line);border-radius:9px;margin:16px 0}table{border-collapse:collapse;width:100%;text-align:left}th{background:#eaf0fb;font-size:13px;color:#304366;white-space:nowrap;position:sticky;top:0}th,td{padding:13px;vertical-align:top;border-bottom:1px solid var(--line)}td{min-width:130px;max-width:330px;overflow-wrap:anywhere}td:first-child{min-width:180px}tr:nth-child(even){background:#fafcff}tr:last-child td{border-bottom:0}small{display:block;color:var(--muted);font-size:12px;margin-top:5px}.note{background:#eff4ff;border-left:4px solid #4162e5;padding:15px 18px;border-radius:6px;margin:14px 0}.warn{background:#fff4e9;border-left-color:#de9329}.danger{color:#a72a35}.issue{border:1px solid var(--line);border-left:5px solid #e2a73b;border-radius:10px;padding:18px;margin:18px 0}.issue.P0{border-left-color:#c64754}.issue.P2{border-left-color:#4c79b9}.issue td:first-child{min-width:105px;max-width:105px}.issue td:nth-child(2){width:100%;max-width:none}.badge{display:inline-block;padding:2px 8px;border-radius:5px;background:#fff0d3;color:#875a14;font-size:12px}.badge.P0{background:#ffe5e9;color:#9b2532}.badge.P2{background:#e4efff;color:#24599e}.sources{font-size:12px;overflow-wrap:anywhere}.hidden,[hidden]{display:none!important}.footer{color:var(--muted);font-size:12px;padding:18px 0}.status{font-size:12px;color:var(--muted)}.cols{display:grid;grid-template-columns:1fr 1fr;gap:18px}.cols section{min-width:0}@media(max-width:800px){h1{font-size:24px}.kpis{grid-template-columns:repeat(3,1fr)}main{padding:0 12px}.panel{padding:16px}.cols{grid-template-columns:1fr}nav{position:static}}@media print{nav,.toolbar{display:none}body{background:white}.panel{break-inside:avoid;border:0}.tab-content[hidden]{display:block!important}.table-wrap{overflow:visible}th{position:static}td{font-size:9px;min-width:0!important;padding:5px}header{background:white;color:black;padding:12px}header p{color:black}.kpis{grid-template-columns:repeat(6,1fr)}}
</style></head><body><header><h1>AIMaster 프로그램 분석·추가 개발 제안서</h1><p>2026-10-19 · 83개 프로그램 · 근거 중심 결함 검토와 창의적 연계/병합 제안</p><p>분석·제안 완료 / 제품 수정·DB 변경·유료 실행·운영 배포 없음</p><a href="REPORT.md">상세 Markdown</a> · <a href="AIMaster_분석_개발제안_20261019.xlsx">엑셀 전체 표</a> · <a href="programs.csv">프로그램 CSV</a> · <a href="ROADMAP.md">CLI 실행표</a></header><main>
'''
html_doc+='<div class="kpis">'+''.join(f'<div class="kpi"><b>{n}</b><span>{h(label)}</span></div>' for n,label in [(83,'등록 프로그램'),(14,'현재 카테고리'),(19,'확인/조건부 문제'),(12,'병합 후보'),(18,'창의적 아이디어'),(0,'제품 유료 호출')])+'</div>'
html_doc+='<nav aria-label="보고서 섹션">'+''.join(f'<button type="button" role="tab" aria-selected="{str(i==0).lower()}" data-tab="{key}">{label}</button>' for i,(key,label) in enumerate([('summary','종합 판단'),('programs','카테고리·프로그램83'),('issues','문제·검증19'),('merges','병합 후보12'),('ideas','아이디어18'),('roadmap','CLI 협업')]))+'</nav>'
html_doc+='<section id="summary" class="tab-content panel"><h2>안정성을 먼저, 연계는 작게 시작</h2><p>공통 접근 정책·직접 API·회원 보관·발송 정합성을 먼저 정비한 뒤, 기존 엔진의 장점을 새 독립 프로그램에 조합하는 순서를 권합니다.</p>'+critical_html+'<div class="cols"><section><h3>확인한 범위</h3><p>운영 카탈로그83개·DB테이블202개 정책 메타데이터, 고정 소스 '+h(base[:8])+', 소스5,193파일 정적 패턴 점검·문서615개 색인. 위험 경로는 코드와 DB를 추가 대조했습니다.</p><p>HEAD83주소는 리디렉트 후200. 빈 입력 API2개·공용 사본 검사·오디오 lint 재현을 실행했습니다. 회원 데이터 원문/키는 읽지 않았습니다.</p></section><section><h3>확인하지 않은 범위</h3><p>전 코드를 사람 눈으로 한 줄씩 심층 검토하거나83개 기능을 실행한 것은 아닙니다. 전체 빌드/의존성 취약점 최신audit/회원E2E/실제 생성·발행·발송·렌더는 미검수입니다.</p><p>HTTP200·RLS켜짐·디자인완료는 기능/보안 보증이 아닙니다. 운영 피해 발생 여부를 시험하지 않았습니다.</p></section></div><div class="note warn">P0 DB 정책은 실제 허용 설정 확인입니다. 이번에는 수정하지 않았습니다. DB 변경 승인과 역할별 검증이 필요합니다. 기존 운영을 안전하게 바꾸는 작업과 아이디어 구현을 구분하세요.</div><h3>현재 분류14개 → 목적 중심 보조분야11개(제안)</h3>'+cat_html+'<p>현재 카테고리/상품은 보존합니다. 새 분류는 사용자 목표·산출물·필요한 계정에 맞춘 탐색 태그입니다.</p></section>'
html_doc+='<section id="programs" class="tab-content panel" hidden><h2>프로그램별 분석 · 83개</h2><p class="muted">검증 과제에는 구조적 제약·아이디어도 포함합니다. 확인된 문제는 F-ID로 연결했습니다. 버전은 수집 시점 스냅샷이며 착수 시 최신 코드/DB를 다시 확인하세요.</p><div class="toolbar"><input id="program-search" aria-label="프로그램 검색" placeholder="이름·기능·슬러그·이슈ID 검색"><select id="category-filter" aria-label="카테고리 필터"><option value="">전체14개 카테고리</option>'+''.join('<option>'+h(c)+'</option>' for c in counts)+'</select><select id="active-filter" aria-label="활성 필터"><option value="">활성·비활성 모두</option><option value="active">활성81개</option><option value="inactive">비활성2개</option></select><button id="reset">필터 초기화</button></div><p id="row-count" class="status">83 / 83개 표시</p>'+program_html+'</section>'
html_doc+='<section id="issues" class="tab-content panel" hidden><h2>확인 문제·조건부 위험 · 19건</h2><p>우선순위와 확인 수준을 함께 보세요. 각 항목은 현재 미수정이며 실제 침해/피해 발생을 의미하지 않습니다.</p><div class="toolbar"><input id="issue-search" aria-label="문제 검색" placeholder="문제·프로그램·ID 검색"><select id="priority-filter" aria-label="우선순위 필터"><option value="">모든 우선순위</option><option>P0</option><option>P1</option><option>P2</option></select></div>'+issue_html+'</section>'
html_doc+='<section id="merges" class="tab-content panel" hidden><h2>장점을 모은 병합 프로그램 · 12개 후보</h2><div class="note">기존 상품을 폐지하거나 서로의 내부API를 호출하는 구조가 아닙니다. 기준 엔진의 검증된 사본을 새 독립 앱에 구성합니다. 회원 본인 키·자료 격리·원생성일30일 규칙을 유지합니다. 전체 복합 기능보다 핵심2~3기능의 MVP부터 선택하세요.</div>'+merge_html+'</section>'
html_doc+='<section id="ideas" class="tab-content panel" hidden><h2>열린 사고의 확장 아이디어 · 18개</h2><p>전부 제안·미구현이며 신규 권한/비용/DB 변경 승인을 의미하지 않습니다.</p>'+idea_html+'</section>'
html_doc+='<section id="roadmap" class="tab-content panel" hidden><h2>다른 CLI와 이어갈 실행 순서</h2><p>Codex: 기획/설계/오류 재현/검수. Claude Code: 복잡한 구현/기능/리팩터링. UI·Antigravity 세부 담당은 주인님 배정. 실제로 작업을 보내거나 담당을 확정하지 않았습니다.</p>'+road_html+'<div class="note warn">같은 프로그램 동시 작업 금지. 차이가 나면 최신 지시·diff·검증 근거를 비교하고, 기존 작업 대체/판단 어려움은 원본 보존 후 주인님 확인. 이 보고서 기록 자체는 구현 승인이나 운영 완료 증거가 아닙니다.</div><p>다음 시작: AGENTS → ACTIVE_WORK/HANDOFF → 프로그램 문서 → F/W-ID 근거 재검증 → 개발 플랜 → 필요한 승인 → 구현/무료회귀 → 필요한 배포 → 완료·미완료 인수인계.</p></section>'
html_doc+='<div class="footer">분석 소스 '+h(base)+' · <a href="summary.json">검수 범위</a> · <a href="http-evidence.json">HTTP 근거</a> · <a href="db-metadata.json">DB정책 근거</a> · <a href="REPORT.md">공식 문서/상세 근거</a><br>본 보고서는 로컬 문서이며 외부 서버·CDN·분석 도구에 사용자 입력을 보내지 않습니다.</div></main><script>const tabs=[...document.querySelectorAll("[data-tab]")];for(const b of tabs)b.addEventListener("click",()=>{for(const x of tabs)x.setAttribute("aria-selected",String(x===b));for(const x of document.querySelectorAll(".tab-content"))x.hidden=x.id!==b.dataset.tab;window.scrollTo(0,0)});const q=document.querySelector("#program-search"),c=document.querySelector("#category-filter"),s=document.querySelector("#active-filter");function filter(){let n=0;for(const r of document.querySelectorAll("#programs tbody tr")){const show=r.dataset.search.toLowerCase().includes(q.value.toLowerCase())&&(!c.value||r.dataset.category===c.value)&&(!s.value||r.dataset.state===s.value);r.hidden=!show;n+=show}document.querySelector("#row-count").textContent=n+" / 83개 표시"}for(const x of [q,c,s]){x.addEventListener("input",filter);x.addEventListener("change",filter)}document.querySelector("#reset").addEventListener("click",()=>{q.value=c.value=s.value="";filter()});const iq=document.querySelector("#issue-search"),ip=document.querySelector("#priority-filter");function fi(){for(const r of document.querySelectorAll("article.issue"))r.hidden=!(r.dataset.search.toLowerCase().includes(iq.value.toLowerCase())&&(!ip.value||r.classList.contains(ip.value)))}iq.addEventListener("input",fi);ip.addEventListener("change",fi);</script></body></html>'
(OUT/'REPORT.html').write_text(html_doc,encoding='utf-8')

# Workbook: real Excel filters, frozen rows, wrapped Korean text, hyperlinks.
wb=Workbook();wb.remove(wb.active)
def sheet(name,headers,rows,widths):
 ws=wb.create_sheet(name);ws.append(headers)
 for row in rows:ws.append([str(v) if isinstance(v,(dict,list)) else v for v in row])
 ws.freeze_panes='A2';ws.sheet_view.zoomScale=80;ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='203D68');cell.alignment=Alignment(vertical='center',wrap_text=True)
 ws.row_dimensions[1].height=30
 for row in ws.iter_rows(min_row=2):
  ws.row_dimensions[row[0].row].height=90 if name not in ['요약','카테고리'] else 48
  for cell in row:
   cell.font=Font(name='맑은 고딕',size=10);cell.alignment=Alignment(vertical='top',wrap_text=True)
   if row[0].row%2==0:cell.fill=PatternFill('solid',fgColor='F1F5FC')
   if isinstance(cell.value,str) and cell.value.startswith(('https://','http://')):cell.hyperlink=cell.value;cell.font=Font(name='맑은 고딕',color='244AC1',underline='single',size=10)
   if isinstance(cell.value,str) and cell.value in ['P0','P1','P2']:
    cell.fill=PatternFill('solid',fgColor={'P0':'FFE0E6','P1':'FFF0CE','P2':'DFEBFF'}[cell.value])
 for i,width in enumerate(widths,1):ws.column_dimensions[ws.cell(1,i).column_letter].width=width
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:1'
 return ws
sheet('요약',['항목','결과·범위'],[['등록','83개 / 활성81·비활성2 / 현재카테고리14'],['문제','19건: 운영설정/코드/조건부/기존기록 구분, 미수정'],['병합·아이디어','병합12·확장18·협업작업12, 전부제안·미배정'],['분석기준',base],['방법','5193개 소스파일 정적점검·615개 문서색인·핵심경로 추가대조'],['실행','83 HEAD·빈입력API2·공용모듈/오디오lint 재현'],['미실행','전체빌드/audit·회원E2E·유료AI·수집·게시·발송·렌더'],['최우선','P0 DB정책/F-001과 eBook API/F-002~003'],['원칙','본인키·회원격리·생성일30일·독립배포·기존작업보존'],['라이브','프로그램 시트의 운영URL 사용. 링크200은 기능검증아님']], [24,115])
sheet('카테고리',['현재분류','개수','활성','비활성','보조분야안'],[[c,n,sum(p['is_active'] for p in programs if p['category']==c),sum(not p['is_active'] for p in programs if p['category']==c),family[c]] for c,n in counts.items()],[26,12,12,12,40])
sheet('프로그램83',program_headers,program_rows,[22,24,34,30,12,14,14,45,52,48,48,20,25,48,48,18,42])
sheet('문제점19',['ID','우선','문제','대상','확인수준','관찰','영향·한계','개선방향','검수기준','승인','근거'],[[x['id'],x['priority'],x['title'],','.join(x['targets']),x['status'],x['observation'],x['impact'],x['fix'],x['verify'],x['approval'],'\n'.join(x['refs'])] for x in ISSUES],[12,10,45,32,30,85,70,85,85,52,65])
sheet('병합제안12',['ID','이름','활용프로그램','가치','가져올장점','줄일단점','첫MVP','위험·선행','규모','검수기준','추천단계'],[[m['id'],m['name'],', '.join(m['sources']),m['value'],m['keep'],m['remove'],m['mvp'],m['risk'],m['effort'],m['gate'],m['phase']] for m in MERGES],[12,34,66,60,65,55,76,78,10,75,14])
sheet('아이디어18',['ID','아이디어','분야','가치·동작','설계경계','완료기준'],IDEAS,[12,42,30,95,90,70])
sheet('협업로드맵12',['ID','우선','작업','이슈','선행','권장역할','승인','완료기준'],ROADMAP,[12,10,44,30,45,90,65,90])
wb.save(OUT/'AIMaster_분석_개발제안_20261019.xlsx')

readme=f'''# 2026-10-19 AIMaster 프로그램 분석 패키지

주인님 요청: 현재 프로그램을 정밀 분석해 카테고리별 버그/로직/기능 개선/연계/병합 개발 제안서를 표로 작성. 이번은 분석·문서 작업이며 구현/DB/키/운영 배포 없음.

- **먼저 보기:** [검색·필터 HTML](REPORT.html), [엑셀](AIMaster_분석_개발제안_20261019.xlsx)
- [전체 근거·프로그램별83개 Markdown](REPORT.md), [협업 실행표](ROADMAP.md), [프로그램 CSV](programs.csv)
- [JSON 재사용 데이터](proposal-data.json), [검수 범위](summary.json), [정적 인벤토리](inventory.json)
- 운영 증거: catalog-db/db-metadata/db-detail-evidence/http-evidence/tooling-evidence JSON 및 shared-module-check.txt.
- 분석 소스 {base}, 프로그램 코드는 변경하지 않음. 고정 소스5193파일 자동정적점검+선별 심층대조이며 전기능실행 인증아님. 개발 중 후속 변경은 착수 전 재확인.
- 보안상 회원 행/키/인증정보 원문을 조회/저장하지 않음. 문서 원문 추출 중간 자료는 git제외 .vercel/analysis-20261019-extracts에만 보관.
- 재생성: python docs/analysis-20261019/build_report.py (openpyxl 필요). 카탈로그/DB 메타 스냅샷은 다시 읽기 전용 조회 후 교체. collect_inventory.py는 현재 HEAD를 고정해 읽음. review_extracts.py는 로컬 중간자료 요약용. probe_public.py는 HEAD83개+빈입력 API2만 실행하며 실제수집/AI/발송없음.
- **다음 첫 단계:** 최신 지침/ACTIVE → F-001 운영 정책 근거 재검증 → 주인님께 DB수정 플랜과 승인범위 확정. W-01~12는 미배정/미구현이고 병합안을 임의 구현하지 않음.
- 완료:83개 분류/정적 조사·19개 문제/조건부위험·12개 병합·18개 아이디어·12개 협업단위·표/근거/인수인계 작성.
- 미완료:문제 수정/DB정비·전체 회원E2E/audit/렌더·신규프로그램·유료 검증. 실제 외부 게시/발송/회원자료 조회0.
'''
(OUT/'README.md').write_text(readme,encoding='utf-8')
from module_design import integrate
integrate(OUT)
print(json.dumps({'programs':len(programs),'issues':len(ISSUES),'merges':len(MERGES),'ideas':len(IDEAS),'families':dict(families),'artifacts':['REPORT.html','REPORT.md','AIMaster_분석_개발제안_20261019.xlsx','programs.csv','proposal-data.json','ROADMAP.md']},ensure_ascii=False))
